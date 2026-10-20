from datetime import UTC, datetime
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)

from enterprise_document_rag.ingestion import (
    build_local_manifest,
)
from enterprise_document_rag.processing import (
    process_local_manifest,
)

GENERATED_AT = datetime(
    2026,
    10,
    20,
    12,
    0,
    tzinfo=UTC,
)


def create_pdf(
    path: Path,
) -> None:
    """Create a small PDF containing extractable synthetic text."""
    writer = PdfWriter()
    page = writer.add_blank_page(
        width=612,
        height=792,
    )

    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_reference = writer._add_object(font)

    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {
                    NameObject("/F1"): font_reference,
                }
            )
        }
    )

    content = DecodedStreamObject()
    content.set_data(b"BT\n/F1 12 Tf\n72 720 Td\n(Synthetic PDF security policy.) Tj\nET\n")
    page[NameObject("/Contents")] = writer._add_object(content)

    with path.open("wb") as destination:
        writer.write(destination)


def create_docx(
    path: Path,
) -> None:
    """Create a small DOCX containing a heading and control table."""
    document = DocxDocument()
    document.add_heading(
        "Synthetic Access Control Standard",
        level=1,
    )

    table = document.add_table(
        rows=2,
        cols=2,
    )
    table.cell(0, 0).text = "Control"
    table.cell(0, 1).text = "Owner"
    table.cell(1, 0).text = "Privileged access review"
    table.cell(1, 1).text = "Information Security"

    document.save(str(path))


def test_processes_pdf_and_docx_through_complete_local_pipeline(
    tmp_path: Path,
) -> None:
    source_root = tmp_path / "source"
    source_root.mkdir()

    create_pdf(
        source_root / "security-policy.pdf",
    )
    create_docx(
        source_root / "access-control-standard.docx",
    )

    manifest = build_local_manifest(
        source_root,
        source_id="synthetic-office-documents",
        access_groups=(
            "engineering",
            "security",
        ),
        generated_at=GENERATED_AT,
    )

    batch = process_local_manifest(
        source_root,
        manifest,
        generated_at=GENERATED_AT,
    )

    assert manifest.document_count == 2
    assert manifest.failure_count == 0
    assert batch.document_count == 2
    assert batch.failure_count == 0
    assert batch.ingestion_failure_count == 0

    documents = {document.source.filename: document for document in batch.documents}

    pdf = documents["security-policy.pdf"]
    assert pdf.processor.name == "pypdf-parser"
    assert pdf.source.access_groups == (
        "engineering",
        "security",
    )
    assert len(pdf.segments) == 1
    assert pdf.segments[0].text == "Synthetic PDF security policy."
    assert pdf.segments[0].location.page_number == 1

    docx = documents["access-control-standard.docx"]
    assert docx.processor.name == "python-docx-parser"
    assert docx.source.access_groups == (
        "engineering",
        "security",
    )
    assert tuple(segment.text for segment in docx.segments) == (
        "Synthetic Access Control Standard",
        "Control | Owner",
        "Privileged access review | Information Security",
    )
