from datetime import UTC, datetime
from pathlib import Path
from typing import NoReturn
from zipfile import BadZipFile

import pytest
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.text.paragraph import Paragraph

import enterprise_document_rag.processing.docx_parser as docx_parser
from enterprise_document_rag.ingestion.local import (
    build_local_metadata,
)
from enterprise_document_rag.ingestion.models import (
    DocumentMetadata,
)
from enterprise_document_rag.processing.docx_parser import (
    parse_docx,
)
from enterprise_document_rag.processing.models import (
    SegmentKind,
)

PROCESSED_AT = datetime(
    2026,
    1,
    1,
    12,
    0,
    tzinfo=UTC,
)


def create_docx_source(
    root: Path,
    *,
    filename: str = "policy.docx",
) -> tuple[Path, DocumentMetadata]:
    path = root / filename
    document = Document()
    document.add_paragraph("Synthetic policy.")
    document.save(str(path))
    metadata = build_local_metadata(
        path,
        root,
        source_id="test-docx-source",
    )
    return path, metadata


def create_structured_docx_source(
    root: Path,
) -> tuple[Path, DocumentMetadata]:
    path = root / "structured.docx"
    document = Document()

    heading = document.add_paragraph("  Access Control Policy  ")
    heading.style = "Title"

    document.add_paragraph("   ")

    table = document.add_table(
        rows=3,
        cols=2,
    )
    table.cell(
        0,
        0,
    ).text = "Control"
    table.cell(
        0,
        1,
    ).text = "Owner"
    table.cell(
        1,
        0,
    ).text = "MFA"
    table.cell(
        1,
        1,
    ).text = "Security\nTeam"
    table.cell(
        2,
        0,
    ).text = ""
    table.cell(
        2,
        1,
    ).text = "   "

    document.add_paragraph("Cafe\u0301 approval.")
    document.save(str(path))

    metadata = build_local_metadata(
        path,
        root,
        source_id="test-docx-source",
    )
    return path, metadata


def test_extracts_paragraphs_and_tables_in_document_order(
    tmp_path: Path,
) -> None:
    path, metadata = create_structured_docx_source(tmp_path)

    processed = parse_docx(
        path,
        metadata,
        processed_at=PROCESSED_AT,
    )

    assert processed.source == metadata
    assert processed.processor.name == "python-docx-parser"
    assert processed.processor.version == "1.0"
    assert processed.processed_at == PROCESSED_AT
    assert len(processed.segments) == 4

    heading, header_row, data_row, closing = processed.segments

    assert heading.ordinal == 0
    assert heading.kind is SegmentKind.TEXT
    assert heading.text == "Access Control Policy"
    assert heading.location.section == "Title"
    assert heading.attributes == {
        "block_index": "1",
        "paragraph_index": "1",
        "style": "Title",
    }

    assert header_row.ordinal == 1
    assert header_row.kind is SegmentKind.TABLE_ROW
    assert header_row.text == "Control | Owner"
    assert header_row.location.row_number == 1
    assert header_row.attributes == {
        "block_index": "3",
        "table_index": "1",
        "row_index": "1",
        "cell_count": "2",
    }

    assert data_row.ordinal == 2
    assert data_row.text == "MFA | Security\nTeam"
    assert data_row.location.row_number == 2

    assert closing.ordinal == 3
    assert closing.kind is SegmentKind.TEXT
    assert closing.text == "Café approval."
    assert closing.location.section == "Normal"
    assert closing.attributes == {
        "block_index": "4",
        "paragraph_index": "3",
        "style": "Normal",
    }

    assert processed.attributes == {
        "paragraph_count": "3",
        "table_count": "1",
        "table_row_count": "2",
        "empty_block_count": "2",
        "normalization": "unicode-nfc",
    }


def test_builds_stable_segment_identities(
    tmp_path: Path,
) -> None:
    path, metadata = create_structured_docx_source(tmp_path)

    first = parse_docx(
        path,
        metadata,
        processed_at=PROCESSED_AT,
    )
    second = parse_docx(
        path,
        metadata,
        processed_at=PROCESSED_AT,
    )

    assert tuple(segment.segment_id for segment in first.segments) == tuple(
        segment.segment_id for segment in second.segments
    )


def test_repeated_table_text_has_distinct_row_identity(
    tmp_path: Path,
) -> None:
    path = tmp_path / "repeated.docx"
    document = Document()
    table = document.add_table(
        rows=2,
        cols=1,
    )
    table.cell(
        0,
        0,
    ).text = "Repeated"
    table.cell(
        1,
        0,
    ).text = "Repeated"
    document.save(str(path))
    metadata = build_local_metadata(
        path,
        tmp_path,
        source_id="test-docx-source",
    )

    processed = parse_docx(
        path,
        metadata,
        processed_at=PROCESSED_AT,
    )

    assert processed.segments[0].segment_id != processed.segments[1].segment_id


def test_uses_current_utc_time_by_default(
    tmp_path: Path,
) -> None:
    path, metadata = create_docx_source(tmp_path)

    before = datetime.now(UTC)
    processed = parse_docx(
        path,
        metadata,
    )
    after = datetime.now(UTC)

    assert before <= processed.processed_at <= after
    assert processed.processed_at.tzinfo is not None
    assert processed.processed_at.utcoffset() is not None


def test_handles_paragraph_without_style(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path, metadata = create_docx_source(tmp_path)

    monkeypatch.setattr(
        Paragraph,
        "style",
        property(lambda self: None),
    )

    processed = parse_docx(
        path,
        metadata,
        processed_at=PROCESSED_AT,
    )

    assert processed.segments[0].text == "Synthetic policy."
    assert processed.segments[0].location.section is None
    assert "style" not in processed.segments[0].attributes


def test_rejects_non_docx_metadata(
    tmp_path: Path,
) -> None:
    path = tmp_path / "policy.txt"
    path.write_text(
        "Synthetic policy.",
        encoding="utf-8",
    )
    metadata = build_local_metadata(
        path,
        tmp_path,
        source_id="test-source",
    )

    with pytest.raises(
        ValueError,
        match="requires DOCX document metadata",
    ):
        parse_docx(
            path,
            metadata,
            processed_at=PROCESSED_AT,
        )


def test_rejects_non_file_document(
    tmp_path: Path,
) -> None:
    path, metadata = create_docx_source(tmp_path)
    path.unlink()
    path.mkdir()

    with pytest.raises(
        ValueError,
        match="not a regular file",
    ):
        parse_docx(
            path,
            metadata,
            processed_at=PROCESSED_AT,
        )


def test_rejects_symbolic_link(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path, metadata = create_docx_source(tmp_path)

    original_is_symlink = Path.is_symlink

    def is_symlink(
        candidate: Path,
    ) -> bool:
        if candidate == path:
            return True
        return original_is_symlink(candidate)

    monkeypatch.setattr(
        Path,
        "is_symlink",
        is_symlink,
    )

    with pytest.raises(
        ValueError,
        match="not a regular file",
    ):
        parse_docx(
            path,
            metadata,
            processed_at=PROCESSED_AT,
        )


def test_rejects_filename_mismatch(
    tmp_path: Path,
) -> None:
    path, metadata = create_docx_source(tmp_path)
    renamed = tmp_path / "renamed.docx"
    renamed.write_bytes(path.read_bytes())

    with pytest.raises(
        ValueError,
        match="filename does not match",
    ):
        parse_docx(
            renamed,
            metadata,
            processed_at=PROCESSED_AT,
        )


def test_rejects_content_hash_mismatch(
    tmp_path: Path,
) -> None:
    path, metadata = create_docx_source(tmp_path)
    path.write_bytes(b"changed DOCX content")

    with pytest.raises(
        ValueError,
        match="content does not match",
    ):
        parse_docx(
            path,
            metadata,
            processed_at=PROCESSED_AT,
        )


@pytest.mark.parametrize(
    "error",
    [
        BadZipFile("synthetic bad ZIP"),
        KeyError("synthetic missing package part"),
        OSError("synthetic read failure"),
        PackageNotFoundError("synthetic package failure"),
        ValueError("synthetic value failure"),
    ],
)
def test_rejects_unreadable_docx(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    error: Exception,
) -> None:
    path, metadata = create_docx_source(tmp_path)

    def fail_document(
        source: str,
    ) -> NoReturn:
        del source
        raise error

    monkeypatch.setattr(
        docx_parser,
        "Document",
        fail_document,
    )

    with pytest.raises(
        ValueError,
        match="DOCX document could not be read",
    ):
        parse_docx(
            path,
            metadata,
            processed_at=PROCESSED_AT,
        )


def test_rejects_real_malformed_docx(
    tmp_path: Path,
) -> None:
    path = tmp_path / "malformed.docx"
    path.write_bytes(b"not a ZIP package")
    metadata = build_local_metadata(
        path,
        tmp_path,
        source_id="test-docx-source",
    )

    with pytest.raises(
        ValueError,
        match="DOCX document could not be read",
    ):
        parse_docx(
            path,
            metadata,
            processed_at=PROCESSED_AT,
        )


def test_requires_processable_content(
    tmp_path: Path,
) -> None:
    path = tmp_path / "empty.docx"
    document = Document()
    document.add_paragraph("   ")
    table = document.add_table(
        rows=1,
        cols=2,
    )
    table.cell(
        0,
        0,
    ).text = ""
    table.cell(
        0,
        1,
    ).text = "   "
    document.save(str(path))
    metadata = build_local_metadata(
        path,
        tmp_path,
        source_id="test-docx-source",
    )

    with pytest.raises(
        ValueError,
        match="contains no processable content",
    ):
        parse_docx(
            path,
            metadata,
            processed_at=PROCESSED_AT,
        )
